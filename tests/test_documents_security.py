"""Gate: negação entre empresas (inclusive download/export/jobs/cópia/sync), cópia sem
IDs fiscais, XML/ZIP malicioso bloqueado, XML download autenticado, exportação integral."""
from __future__ import annotations

import io
import uuid
import zipfile
from datetime import datetime, timedelta, timezone

from conftest import SP

F = "/api/fiscal/v1/companies"


def _xml(world, number=777, toma=None, prest=None):  # noqa: ANN001, ANN201
    from nfse_api.providers.mock_data import cnpj_dv, mock_access_key, nfse_xml

    prest = prest or (cnpj_dv("700000000001"), "Fornecedor Importado")
    key = mock_access_key(prest[0], number, datetime(2026, 10, 1, tzinfo=timezone.utc), str(number))
    xml = nfse_xml(key=key, number=number, prest=prest, toma=toma or (world["a_cnpj"], "Empresa A"), interm=None,
                   description="Serviço importado", v_serv="300.00", issued_at=datetime(2026, 10, 1, 10, tzinfo=SP),
                   competence_day="2026-10-01")
    return key, xml


def _import(stack, world, files, h="ha", company="a"):  # noqa: ANN001, ANN201
    r = stack.fiscal.post(f"{F}/{world[company]}/documents/import", headers=world[h],
                          files=[("files", (n, d, "application/octet-stream")) for n, d in files])
    assert r.status_code == 202, r.text
    stack.worker().run_jobs()
    return stack.fiscal.get(f"{F}/{world[company]}/jobs/{r.json()['job_id']}", headers=world[h]).json()


def test_import_review_and_authenticated_xml_download(stack, world):
    key, xml = _xml(world)
    job = _import(stack, world, [("nota.xml", xml)])
    assert job["status"] == "SUCCEEDED" and job["result"]["created"] == 1
    d = stack.fiscal.get(f"{F}/{world['a']}/documents", headers=world["ha"], params={"query": key}).json()["items"][0]
    assert d["direction"] == "RECEIVED" and d["source"] == "IMPORT" and d["simulated"] is False
    url = f"{F}/{world['a']}/documents/{d['id']}/xml"
    assert stack.fiscal.get(url).status_code == 401
    assert stack.fiscal.get(url, headers=world["hb"]).status_code == 403
    assert stack.fiscal.get(f"{F}/{world['b']}/documents/{d['id']}/xml", headers=world["hb"]).status_code == 404
    ok = stack.fiscal.get(url, headers=world["ha"])
    assert ok.status_code == 200 and ok.content == xml and "attachment" in ok.headers["content-disposition"]
    # Conferência: autor/data, controle de versão, somente leitor não pode.
    rv = f"{F}/{world['a']}/documents/{d['id']}/review"
    assert stack.fiscal.patch(rv, headers=world["hr"], json={"review_status": "REVIEWED", "expected_version": 1}).status_code == 403
    r = stack.fiscal.patch(rv, headers=world["ha"], json={"review_status": "DIVERGENT", "note": "ISS", "expected_version": 1})
    assert r.status_code == 200 and r.json()["reviewed_by"] == world["ua"] and r.json()["version"] == 2
    assert stack.fiscal.patch(rv, headers=world["ha"], json={"review_status": "REVIEWED", "expected_version": 1}).status_code == 409
    assert ok.content == stack.fiscal.get(url, headers=world["ha"]).content  # XML original intocado
    again = _import(stack, world, [("dup.xml", xml.replace(b"><", b">\n<", 3))])
    assert again["result"]["duplicates"] == 1 and again["errors"][0]["code"] == "DUPLICATE_IGNORED"


def test_malicious_xml_and_zip_blocked(stack, world):
    xxe = b'<?xml version="1.0"?><!DOCTYPE r [<!ENTITY x SYSTEM "file:///etc/passwd">]><NFSe>&x;</NFSe>'
    bomb = b'<?xml version="1.0"?><!DOCTYPE l [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;">]><NFSe>&b;</NFSe>'
    nfe = b'<?xml version="1.0"?><nfeProc xmlns="http://www.portalfiscal.inf.br/nfe"><NFe/></nfeProc>'
    _, good = _xml(world, 801)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("../../etc/evil.xml", good)  # nome com traversal: só vira rótulo sanitizado
        z.writestr("ok/inner.zip", b"PK\x03\x04")
        z.writestr("leia.txt", b"oi")
        z.writestr("bomba.xml", b"<a>" + b" " * 3_000_000 + b"</a>")  # > limite e razão alta
    job = _import(stack, world, [("xxe.xml", xxe), ("bomb.xml", bomb), ("nfe.xml", nfe), ("lote.zip", buf.getvalue())])
    codes = {(e["file"], e["code"]) for e in job["errors"]}
    assert ("xxe.xml", "XML_FORBIDDEN_CONSTRUCT") in codes
    assert ("bomb.xml", "XML_FORBIDDEN_CONSTRUCT") in codes
    assert ("nfe.xml", "NFE_NOT_SUPPORTED") in codes
    assert ("lote.zip/inner.zip", "NESTED_ARCHIVE") in codes
    assert ("lote.zip/leia.txt", "UNSUPPORTED_ENTRY") in codes
    assert any(f == "lote.zip/bomba.xml" and c in ("XML_TOO_LARGE", "ZIP_RATIO_EXCEEDED") for f, c in codes)
    assert job["result"]["created"] == 1 and job["status"] == "PARTIAL"  # evil.xml (conteúdo válido) importado
    assert all(".." not in (e["file"] or "") for e in job["errors"])
    # Muitas entradas: ZIP inteiro recusado.
    many = io.BytesIO()
    with zipfile.ZipFile(many, "w") as z:
        for i in range(501):
            z.writestr(f"{i}.xml", b"<a/>")
    job = _import(stack, world, [("muitos.zip", many.getvalue())])
    assert job["errors"][0]["code"] == "ZIP_TOO_MANY_ENTRIES"
    r = stack.fiscal.post(f"{F}/{world['a']}/documents/import", headers=world["ha"],
                          files=[("files", ("x.pdf", b"%PDF-1.4", "application/pdf"))])
    assert r.status_code == 422 and r.json()["code"] == "UNSUPPORTED_FILE"


def test_import_respects_company_and_full_cnpj(stack, world):
    from nfse_api.providers.mock_data import cnpj_dv

    branch = cnpj_dv(world["a_cnpj"][:8] + "0002")
    _, other_branch = _xml(world, 901, toma=(branch, "Filial"))
    _, of_b = _xml(world, 902, toma=(world["b_cnpj"], "Empresa B"))
    job = _import(stack, world, [("filial.xml", other_branch), ("b.xml", of_b)])
    codes = sorted(e["code"] for e in job["errors"])
    assert codes == ["NOT_COMPANY_DOCUMENT", "OTHER_BRANCH_SAME_ROOT"] and job["status"] == "FAILED"
    assert stack.fiscal.get(f"{F}/{world['a']}/documents", headers=world["ha"]).json()["total"] == 0


def test_cross_company_denial_everywhere(stack, world):
    a, b = world["a"], world["b"]
    hb = world["hb"]
    _, xml = _xml(world, 950)
    job = _import(stack, world, [("n.xml", xml)])
    doc = stack.fiscal.get(f"{F}/{a}/documents", headers=world["ha"]).json()["items"][0]
    ex = stack.fiscal.post(f"{F}/{a}/exports", headers=world["ha"], json={"format": "xlsx"}).json()
    stack.worker().run_jobs()
    draft = stack.fiscal.post(f"{F}/{a}/drafts", headers=world["ha"], json={
        "customer_id": world["customer"], "competence": "2026-10", "amount": "10.00",
        "service": {"national_code": "010101", "description": "x"}}).json()
    targets = [
        ("get", f"{F}/{a}/documents"), ("get", f"{F}/{a}/documents/{doc['id']}"),
        ("get", f"{F}/{a}/documents/{doc['id']}/xml"), ("get", f"{F}/{a}/jobs/{job['id']}"),
        ("get", f"{F}/{a}/jobs/{ex['job_id']}/download"), ("get", f"{F}/{a}/dashboard"),
        ("get", f"{F}/{a}/sync/status"), ("get", f"{F}/{a}/drafts/{draft['id']}"),
        ("get", f"{F}/{a}/customers/{world['customer']}/last-issued-document"),
        ("get", f"/api/core/v1/companies/{a}/customers"), ("get", f"/api/core/v1/companies/{a}/integration-status"),
    ]
    for method, url in targets:
        client = stack.core if url.startswith("/api/core") else stack.fiscal
        assert getattr(client, method)(url, headers=hb).status_code == 403, url
    posts = [
        (f"{F}/{a}/drafts/from-document", {"source_document_id": doc["id"]}),
        (f"{F}/{a}/exports", {"format": "xlsx"}),
        (f"{F}/{a}/sync/run", None),
        (f"{F}/{a}/batches", {"drafts": [{"id": draft["id"], "expected_version": 1}],
                              "scheduled_at": "2026-12-01T00:01:00-03:00", "timezone": "America/Sao_Paulo"}),
    ]
    for url, body in posts:
        r = stack.fiscal.post(url, headers={**hb, "Idempotency-Key": uuid.uuid4().hex}, json=body)
        assert r.status_code == 403, url
    # IDs da empresa A usados sob a rota da empresa B (onde B é membro): 404, sem vazamento.
    for url in (f"{F}/{b}/documents/{doc['id']}", f"{F}/{b}/jobs/{job['id']}", f"{F}/{b}/jobs/{ex['job_id']}/download",
                f"{F}/{b}/drafts/{draft['id']}"):
        assert stack.fiscal.get(url, headers=hb).status_code == 404, url
    r = stack.fiscal.post(f"{F}/{b}/drafts/from-document", headers={**hb, "Idempotency-Key": uuid.uuid4().hex},
                          json={"source_document_id": doc["id"]})
    assert r.status_code == 404
    # Permissão: leitor não prepara nem agenda nem importa.
    hr = world["hr"]
    assert stack.fiscal.post(f"{F}/{a}/drafts", headers=hr, json={}).status_code in (403, 422)
    assert stack.fiscal.post(f"{F}/{a}/documents/import", headers=hr, files=[("files", ("a.xml", b"<a/>"))]).status_code == 403


def test_copy_last_issued_without_fiscal_ids(stack, world):
    stack.clock.set(datetime(2026, 10, 20, 8, 0, tzinfo=SP))
    w = stack.worker()
    w.refresh_companies(force=True)
    while w.run_sync():
        pass
    last = stack.fiscal.get(f"{F}/{world['a']}/customers/{world['customer']}/last-issued-document", headers=world["ha"])
    assert last.status_code == 200 and last.json()["number"] == "303"
    src = last.json()
    r = stack.fiscal.post(f"{F}/{world['a']}/drafts/from-document", headers={**world["ha"], "Idempotency-Key": uuid.uuid4().hex},
                          json={"source_document_id": src["id"]})
    assert r.status_code == 201
    d = r.json()
    assert d["id"] != src["id"] and d["status"] == "DRAFT" and d["version"] == 1
    assert d["source_document_id"] == src["id"] and d["customer_id"] == world["customer"]
    for forbidden in ("number", "access_key", "dps_id", "protocol", "signature", "scheduled_at", "review_status"):
        assert forbidden not in d
    assert d["current_item_id"] is None and d["competence"] == "2026-10"  # competência atual e revisável
    assert d["amount"] == src["amount"] and d["service"]["description"] == src["description"]
    # Nota recebida não serve de referência.
    rec = stack.fiscal.get(f"{F}/{world['a']}/documents", headers=world["ha"], params={"direction": "RECEIVED"}).json()["items"][0]
    bad = stack.fiscal.post(f"{F}/{world['a']}/drafts/from-document", headers={**world["ha"], "Idempotency-Key": uuid.uuid4().hex},
                            json={"source_document_id": rec["id"]})
    assert bad.status_code == 422 and bad.json()["code"] == "SOURCE_NOT_ISSUED"
    none = stack.fiscal.get(f"{F}/{world['a']}/customers/{uuid.uuid4()}/last-issued-document", headers=world["ha"])
    assert none.status_code == 404


def test_export_full_without_truncation_and_limit(stack, world, tmp_path):
    from openpyxl import load_workbook

    files = []
    for i in range(30):
        prest = (world["a_cnpj"], "Empresa A")
        from nfse_api.providers.mock_data import cnpj_dv

        _, xml = _xml(world, 1000 + i, toma=(cnpj_dv("123456780001"), "=HYPERLINK(\"x\")") if i == 0 else None,
                      prest=prest if i == 0 else None)
        files.append((f"n{i}.xml", xml))
    _import(stack, world, files)
    r = stack.fiscal.post(f"{F}/{world['a']}/exports", headers=world["ha"], json={"format": "xlsx", "page": 2, "page_size": 5})
    assert r.status_code == 202
    jid = r.json()["job_id"]
    nd = stack.fiscal.get(f"{F}/{world['a']}/jobs/{jid}/download", headers=world["ha"])
    assert nd.status_code == 409 and nd.json()["code"] == "JOB_NOT_READY"
    stack.worker().run_jobs()
    job = stack.fiscal.get(f"{F}/{world['a']}/jobs/{jid}", headers=world["ha"]).json()
    assert job["status"] == "SUCCEEDED" and job["result"] == {"rows": 30, "matched": 30}
    dl = stack.fiscal.get(f"{F}/{world['a']}/jobs/{jid}/download", headers=world["ha"])
    assert dl.status_code == 200
    p = tmp_path / "x.xlsx"
    p.write_bytes(dl.content)
    ws = load_workbook(p).active
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) == 31  # cabeçalho + todos (paginação ignorada)
    assert any(str(c).startswith("'=HYPERLINK") for row in rows for c in row if c)
    # Limite: falha explícita, nunca arquivo truncado.
    stack.worker_settings.export_max_rows = 10
    r = stack.fiscal.post(f"{F}/{world['a']}/exports", headers=world["ha"], json={"format": "xlsx"})
    stack.worker().run_jobs()
    job = stack.fiscal.get(f"{F}/{world['a']}/jobs/{r.json()['job_id']}", headers=world["ha"]).json()
    assert job["status"] == "FAILED" and job["errors"][0]["code"] == "EXPORT_TOO_LARGE" and job["download_available"] is False
