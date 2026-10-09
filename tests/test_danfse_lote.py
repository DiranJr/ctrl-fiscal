"""DANFSe (NT SE/CGNFS-e 008 v1.02) gerado a partir do XML oficial, filtros por retenção e download em lote (ZIP)
por job assíncrono: conjunto fixo, além da primeira página, manifesto, resultado parcial, isolamento entre
empresas, ticket de uso único (sem token na URL) e expiração. Governo simulado; nenhuma chamada de rede."""
from __future__ import annotations

import io
import json
import zipfile
from datetime import datetime, timedelta

import pytest
from pypdf import PdfReader

from conftest import SP, cnpj
from test_substitution import _event, _ingest

F = "/api/fiscal/v1/companies"
IBS_NFSE = (b"<IBSCBS><cLocalidadeIncid>1505536</cLocalidadeIncid><xLocalidadeIncid>Parauapebas</xLocalidadeIncid>"
            b"<valores><vBC>900.00</vBC><uf><pIBSUF>0.10</pIBSUF><pAliqEfetUF>0.10</pAliqEfetUF></uf>"
            b"<mun><pIBSMun>0.00</pIBSMun><pAliqEfetMun>0.00</pAliqEfetMun></mun>"
            b"<fed><pCBS>0.90</pCBS><pAliqEfetCBS>0.90</pAliqEfetCBS></fed></valores>"
            b"<totCIBS><vTotNF>812.40</vTotNF><gIBS><vIBSTot>0.90</vIBSTot><gIBSUFTot><vIBSUF>0.90</vIBSUF></gIBSUFTot>"
            b"<gIBSMunTot><vIBSMun>0.00</vIBSMun></gIBSMunTot></gIBS><gCBS><vCBS>8.10</vCBS></gCBS></totCIBS></IBSCBS>")
IBS_DPS = (b"<IBSCBS><finNFSe>0</finNFSe><indFinal>0</indFinal><cIndOp>100301</cIndOp><indDest>0</indDest>"
           b"<valores><trib><gIBSCBS><CST>000</CST><cClassTrib>000001</cClassTrib></gIBSCBS></trib></valores></IBSCBS>")


def _nota(prest: str, toma: str, number: int, *, desc: str = "Servico de teste", tp_ret_iss: str = "1",
          v_issqn: str | None = None, fed: dict | None = None, v_total_ret: str | None = None,
          v_liq: str | None = None, ibs: bool = False, tp_amb: str = "2") -> tuple[str, bytes]:
    from nfse_api.providers.mock_data import mock_access_key, nfse_xml

    when = datetime(2026, 10, 1, 9, 0, tzinfo=SP) + timedelta(minutes=number)
    key = mock_access_key(prest, number, when, f"n{number}")
    xml = nfse_xml(key=key, number=number, prest=(prest, "Prestador Teste"), toma=(toma, "Tomador Teste"),
                   interm=None, description=desc, v_serv="1000.00", issued_at=when, competence_day="2026-10-01",
                   tp_ret_iss=tp_ret_iss, v_issqn=v_issqn, fed=fed, v_total_ret=v_total_ret, v_liq=v_liq)
    if ibs:
        xml = xml.replace(b"<DPS ", IBS_NFSE + b"<DPS ", 1).replace(b"</infDPS>", IBS_DPS + b"</infDPS>", 1)
    if tp_amb != "2":
        xml = xml.replace(b"<tpAmb>2</tpAmb>", f"<tpAmb>{tp_amb}</tpAmb>".encode(), 1)
    return key, xml


def _docs(stack, world, **params) -> dict:  # noqa: ANN001
    r = stack.fiscal.get(f"{F}/{world['a']}/documents", headers=world["ha"], params={"page_size": 100, **params})
    assert r.status_code == 200, r.text
    return {d["access_key"]: d for d in r.json()["items"]}


def _pdf_text(pdf: bytes) -> str:
    r = PdfReader(io.BytesIO(pdf))
    assert len(r.pages) == 1  # NT 008, 2.2: página única obrigatória
    return r.pages[0].extract_text()


# ---------------------------------------------------------------- DANFSe
def test_danfse_from_official_xml_authenticated_and_isolated(stack, world):
    toma = cnpj("445556660001")
    key, xml = _nota(world["a_cnpj"], toma, 1, desc=("Descricao longa " + "palavra " * 245 + "FIM-DA-DESCRICAO")[:2000],
                     tp_ret_iss="2", v_issqn="50.00",
                     fed={"tpRetPisCofins": "3", "vPis": "6.50", "vCofins": "30.00", "vRetCP": "110.00",
                          "vRetIRRF": "15.00", "vRetCSLL": "46.50"},
                     v_total_ret="221.50", v_liq="778.50", ibs=True)
    _ingest(stack, world["a"], world["a_cnpj"], xml)
    d = _docs(stack, world)[key]
    assert d["danfse_available"] is True and d["danfse_watermark"] is None
    url = f"{F}/{world['a']}/documents/{d['id']}/danfse"
    assert stack.fiscal.get(url).status_code == 401
    assert stack.fiscal.get(url, headers=world["hb"]).status_code == 403
    assert stack.fiscal.get(f"{F}/{world['b']}/documents/{d['id']}/danfse", headers=world["hb"]).status_code == 404
    r = stack.fiscal.get(url, headers=world["hr"])  # somente leitura pode visualizar
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert r.headers["content-disposition"].startswith("inline;") and r.headers["cache-control"] == "no-store"
    assert r.headers["x-danfse-origem"] == "gerado-pela-plataforma"
    assert stack.fiscal.get(url, headers=world["ha"], params={"disposition": "attachment"}).headers[
        "content-disposition"].startswith("attachment;")
    meta = PdfReader(io.BytesIO(r.content)).metadata
    assert "gerado pela plataforma" in meta.producer and "008 v1.02" in meta.producer
    text = _pdf_text(r.content)
    for expected in ("DANFSe v2.0", "Documento Auxiliar da NFS-e", "NFS-e SEM VALIDADE JURÍDICA", key,
                     "Retido pelo Tomador", "R$ 50,00", "R$ 110,00", "R$ 15,00", "R$ 46,50",
                     "PIS/COFINS/CSLL Retidos", "R$ 6,50", "R$ 30,00", "R$ 221,50", "R$ 778,50",
                     "000 / 000001", "R$ 812,40", "R$ 8,10", "O DESTINATÁRIO É O PRÓPRIO TOMADOR/ADQUIRENTE",
                     "INTERMEDIÁRIO DA OPERAÇÃO NÃO IDENTIFICADO NA NFS-e", "Totais Aproximados dos Tributos"):
        assert expected in text, expected
    assert "FIM-DA-DESCRICAO" in text  # 2000 caracteres (máximo do leiaute) cabem inteiros na página única
    assert "CTRL Fiscal" not in text  # nada além do XML é impresso (item 2.1)


def test_danfse_code1_cancel_substitution_and_production_header(stack, world):
    toma = cnpj("445556660001")
    k1, x1 = _nota(world["a_cnpj"], toma, 2, fed={"tpRetPisCofins": "1", "vPis": "6.50", "vCofins": "30.00",
                                                   "vRetCSLL": "10.00"}, tp_amb="1")
    k2, x2 = _nota(world["a_cnpj"], toma, 3)
    for x in (x1, x2):
        _ingest(stack, world["a"], world["a_cnpj"], x)
    _ingest(stack, world["a"], world["a_cnpj"], _event(k1, "101101", "<xDesc>Cancelamento de NFS-e</xDesc><cMotivo>1</cMotivo>"
                                                       "<xMotivo>Erro</xMotivo>", world["a_cnpj"]))
    _ingest(stack, world["a"], world["a_cnpj"], _event(
        k2, "105102", f"<xDesc>x</xDesc><cMotivo>01</cMotivo><chSubstituta>{k1}</chSubstituta>", world["a_cnpj"]))
    d = _docs(stack, world)
    assert d[k1]["danfse_watermark"] == "CANCELADA" and d[k2]["danfse_watermark"] == "SUBSTITUÍDA"
    t1 = _pdf_text(stack.fiscal.get(f"{F}/{world['a']}/documents/{d[k1]['id']}/danfse", headers=world["ha"]).content)
    t2 = _pdf_text(stack.fiscal.get(f"{F}/{world['a']}/documents/{d[k2]['id']}/danfse", headers=world["ha"]).content)
    assert "CANCELADA" in t1 and "SUBSTITUÍDA" in t2
    assert "SEM VALIDADE" not in t1  # tpAmb=1: sem a expressão de homologação
    # NT 008 v1.02: modalidade 1 → contribuições retidas = vRetCSLL + vPis + vCofins; PIS/COFINS próprios = 0,00
    assert "R$ 46,50" in t1 and "PIS/COFINS Retidos" in t1
    s = d[k1]["retention_summary"]["contributions"]
    assert s["amount"] == "46.50" and s["consolidated"] is True and s["covered"] == {"pis": True, "cofins": True, "csll": None}


def test_danfse_refused_for_simulated_or_missing_xml(stack, world):
    from ctrl_fiscal.models import Document

    key, xml = _nota(world["a_cnpj"], cnpj("445556660001"), 4)
    _ingest(stack, world["a"], world["a_cnpj"], xml)
    doc_id = _docs(stack, world)[key]["id"]
    with stack.fiscal_sessions() as s, s.begin():
        s.get(Document, doc_id).simulated = True
    r = stack.fiscal.get(f"{F}/{world['a']}/documents/{doc_id}/danfse", headers=world["ha"])
    assert r.status_code == 409 and r.json()["code"] == "DANFSE_SIMULATED"
    with stack.fiscal_sessions() as s, s.begin():
        d = s.get(Document, doc_id)
        d.simulated, d.xml_blob_id = False, None
    r = stack.fiscal.get(f"{F}/{world['a']}/documents/{doc_id}/danfse", headers=world["ha"])
    assert r.status_code == 409 and r.json()["code"] == "XML_UNAVAILABLE"


# ---------------------------------------------------------------- filtros por retenção
@pytest.fixture
def notas(stack, world):  # noqa: ANN001, ANN201
    toma, p = cnpj("445556660001"), world["a_cnpj"]
    spec = {
        "iss_inss": dict(tp_ret_iss="2", v_issqn="50.00", fed={"vRetCP": "110.00"}, v_total_ret="160.00"),
        "sem": dict(),
        "irrf": dict(fed={"vRetIRRF": "15.00"}, v_total_ret="15.00"),
        "contrib": dict(fed={"tpRetPisCofins": "3", "vPis": "6.50", "vCofins": "30.00", "vRetCSLL": "46.50"},
                        v_total_ret="46.50"),
        "informado": dict(fed={"tpRetPisCofins": "2", "vPis": "6.50", "vCofins": "30.00"}),
        "desconhecido": dict(fed={"tpRetPisCofins": "3", "vPis": "6.50", "vCofins": "30.00"}),
    }
    keys = {}
    for n, (name, kw) in enumerate(spec.items(), 10):
        k, x = _nota(p, toma, n, **kw)
        _ingest(stack, world["a"], world["a_cnpj"], x)
        keys[name] = k
    return keys


def _names(stack, world, notas, **params) -> set[str]:  # noqa: ANN001
    found = _docs(stack, world, **params)
    return {n for n, k in notas.items() if k in found}


def test_retention_filters_any_all_range_and_classification(stack, world, notas):
    assert _names(stack, world, notas, retention="WITH") == {"iss_inss", "irrf", "contrib"}
    # imposto informado (vPis/vCofins, código 2) NÃO é retenção
    assert _names(stack, world, notas, retention="WITHOUT") == {"sem", "informado"}
    assert _names(stack, world, notas, retention="UNKNOWN") == {"desconhecido"}  # modalidade retida sem vRetCSLL
    assert _names(stack, world, notas, retention_type="CONTRIB") == {"contrib"}
    assert _names(stack, world, notas, retention_type=["ISS", "INSS"], retention_match="ALL") == {"iss_inss"}
    assert _names(stack, world, notas, retention_type=["ISS", "IRRF"], retention_match="ANY") == {"iss_inss", "irrf"}
    r = stack.fiscal.get(f"{F}/{world['a']}/documents", headers=world["ha"], params={"retention_type": ["ISS", "IRRF"]})
    assert r.status_code == 422 and r.json()["code"] == "RETENTION_MATCH_REQUIRED"
    assert _names(stack, world, notas, withheld_min="20", withheld_max="200") == {"iss_inss", "contrib"}
    assert _names(stack, world, notas, withheld_max="0") == {"sem", "informado"}  # sem retenção = total 0
    r = stack.fiscal.get(f"{F}/{world['a']}/documents", headers=world["ha"], params={"withheld_min": "5", "withheld_max": "1"})
    assert r.status_code == 422
    d = _docs(stack, world)[notas["informado"]]["retention_summary"]
    assert d["classification"] == "WITHOUT" and d["informed_not_withheld"] == {"pis": "6.50", "cofins": "30.00"}
    assert d["contributions"]["amount"] is None


# ---------------------------------------------------------------- download em lote
def _bundle(stack, world, body, h="ha", company="a") -> dict:  # noqa: ANN001
    r = stack.fiscal.post(f"{F}/{world[company]}/documents/bundles", headers=world[h], json=body)
    return r


def _run(stack, world, job_id, h="ha", company="a") -> dict:  # noqa: ANN001
    stack.worker().run_jobs()
    return stack.fiscal.get(f"{F}/{world[company]}/jobs/{job_id}", headers=world[h]).json()


def _zip(stack, world, job_id) -> zipfile.ZipFile:  # noqa: ANN001
    t = stack.fiscal.post(f"{F}/{world['a']}/jobs/{job_id}/download-ticket", headers=world["ha"])
    assert t.status_code == 200, t.text
    body = t.json()
    assert body["method"] == "POST" and body["field"] == "ticket" and "ticket" not in body["action"]
    r = stack.fiscal.post(f"{F}/{world['a']}/jobs/{job_id}/download-file", data={"ticket": body["ticket"]})
    assert r.status_code == 200 and r.headers["content-type"] == "application/zip"
    again = stack.fiscal.post(f"{F}/{world['a']}/jobs/{job_id}/download-file", data={"ticket": body["ticket"]})
    assert again.status_code == 403  # uso único
    return zipfile.ZipFile(io.BytesIO(r.content))


def test_bundle_all_filter_results_beyond_first_page_fixed_set(stack, world, notas):
    page1 = stack.fiscal.get(f"{F}/{world['a']}/documents", headers=world["ha"], params={"page_size": 2}).json()
    assert page1["total"] == 6 and len(page1["items"]) == 2
    r = _bundle(stack, world, {"format": "PDF_XML", "filters": {}})
    assert r.status_code == 202, r.text
    job_id = r.json()["job_id"]
    # conjunto fixado na criação: nota nova depois do pedido não entra
    k, x = _nota(world["a_cnpj"], cnpj("445556660001"), 50)
    _ingest(stack, world["a"], world["a_cnpj"], x)
    job = _run(stack, world, job_id)
    assert job["type"] == "BUNDLE" and job["status"] == "SUCCEEDED", job
    assert job["result"]["documents"] == 6 and job["result"]["pdf_files"] == 6 and job["result"]["xml_files"] == 6
    assert job["download_available"] is True and job["output_expires_at"]
    z = _zip(stack, world, job_id)
    names = z.namelist()
    assert "manifesto.json" in names and "manifesto.csv" in names
    folders = {n.split("/")[0] for n in names if "/" in n}
    assert len(folders) == 6 and not any(k in n for n in names)
    for n in names:
        assert ".." not in n and not n.startswith("/")
    m = json.loads(z.read("manifesto.json"))
    assert m["documents_requested"] == 6 and {e["status"] for e in m["documents"]} == {"OK"}
    assert "não é arquivo baixado do governo" in m["pdf_origin"]
    xml_name = next(n for n in names if n.endswith(".xml"))
    key = xml_name.split("NFSe_")[1][:-4]
    doc = _docs(stack, world)[key]
    assert z.read(xml_name) == stack.fiscal.get(f"{F}/{world['a']}/documents/{doc['id']}/xml", headers=world["ha"]).content


def test_bundle_selection_isolation_partial_failure_and_expiry(stack, world, notas, monkeypatch):
    d = _docs(stack, world)
    ids = [d[notas["sem"]]["id"], d[notas["irrf"]]["id"]]
    # documento de outra empresa: recusado sem revelar existência
    r = _bundle(stack, world, {"format": "XML", "document_ids": ids}, h="hb", company="b")
    assert r.status_code == 422 and r.json()["code"] == "BUNDLE_DOCUMENTS_NOT_FOUND"
    assert _bundle(stack, world, {"format": "XML", "document_ids": ids, "filters": {}}).status_code == 422
    # falha de PDF em um documento preserva o XML e o lote fica PARCIAL
    import ctrl_fiscal.danfse as danfse

    real = danfse.gerar
    bad_key = notas["irrf"]

    def flaky(xml: bytes, marca=None):  # noqa: ANN001, ANN202
        if bad_key.encode() in xml:
            raise RuntimeError("falha simulada")
        return real(xml, marca)

    monkeypatch.setattr(danfse, "gerar", flaky)
    job_id = _bundle(stack, world, {"format": "PDF_XML", "document_ids": ids}).json()["job_id"]
    other = stack.fiscal.get(f"{F}/{world['b']}/jobs/{job_id}", headers=world["hb"])
    assert other.status_code in (403, 404)
    job = _run(stack, world, job_id)
    assert job["status"] == "PARTIAL" and job["result"]["xml_files"] == 2 and job["result"]["pdf_files"] == 1
    assert {e["code"] for e in job["errors"]} == {"DANFSE_FAILED"}
    m = json.loads(_zip(stack, world, job_id).read("manifesto.json"))
    st = {e["access_key"]: e["status"] for e in m["documents"]}
    assert st[bad_key] == "PARTIAL" and st[notas["sem"]] == "OK"
    # ticket de outra empresa / sem ticket
    t = stack.fiscal.post(f"{F}/{world['a']}/jobs/{job_id}/download-ticket", headers=world["ha"]).json()["ticket"]
    assert stack.fiscal.post(f"{F}/{world['b']}/jobs/{job_id}/download-file", data={"ticket": t}).status_code == 403
    assert stack.fiscal.post(f"{F}/{world['a']}/jobs/{job_id}/download-file").status_code == 422
    # expiração: após o prazo, 410 e o worker apaga o arquivo
    stack.clock.advance(hours=25)
    assert stack.fiscal.get(f"{F}/{world['a']}/jobs/{job_id}/download", headers=world["ha"]).status_code == 410
    from ctrl_fiscal.domain.bundles import purge_expired
    from ctrl_fiscal.models import Job

    w = stack.worker()
    assert purge_expired(w.sessions, w.clock, w.settings.export_dir) >= 1
    with stack.fiscal_sessions() as s:
        assert s.get(Job, job_id).output_path is None


def test_bundle_limits_and_empty(stack, world, notas):
    r = _bundle(stack, world, {"format": "PDF", "filters": {"retention": "UNKNOWN", "query": "nao-existe"}})
    assert r.status_code == 422 and r.json()["code"] == "BUNDLE_EMPTY"
    r = _bundle(stack, world, {"format": "PDF", "filters": {"retention_type": ["ISS", "IRRF"]}})
    assert r.status_code == 422 and r.json()["code"] == "RETENTION_MATCH_REQUIRED"
    r = _bundle(stack, world, {"format": "PDF", "filters": {"retention": "WITH"}})
    assert r.status_code == 202
    assert _run(stack, world, r.json()["job_id"])["result"]["documents"] == 3


def test_danfse_max_description_and_complementary_info_single_page(world):
    from ctrl_fiscal.danfse import gerar

    desc = ("Descricao " + "longa " * 400)[:2000]
    key, xml = _nota(world["a_cnpj"] if isinstance(world, dict) else cnpj("112223330001"), cnpj("445556660001"), 60,
                     desc=desc)
    info = ("Informacao complementar " * 100)[:2000]
    xml = xml.replace(b"</cServ>", b"</cServ><infoCompl><xInfComp>" + info.encode() + b"</xInfComp></infoCompl>", 1)
    text = _pdf_text(gerar(xml, "CANCELADA"))
    flat = " ".join(text.split())
    # Descrição e Inf. Complementares no máximo do leiaute (2000 caracteres cada) cabem inteiras em uma página;
    # a linha de totais aproximados (fixa) e a marca d'água continuam presentes.
    assert info.split()[-1] in flat and "longa longa" in flat and "Totais Aproximados dos Tributos" in flat
    assert "CANCELADA" in flat


# ---------------------------------------------------------------- eventos × PDF, ticket, modalidade 1 na exportação
def test_status_change_after_generation_refreshes_pdf_and_blocks_stale_zip(stack, world):
    toma = cnpj("445556660001")
    k, x = _nota(world["a_cnpj"], toma, 70)
    _ingest(stack, world["a"], world["a_cnpj"], x)
    d = _docs(stack, world)[k]
    url = f"{F}/{world['a']}/documents/{d['id']}/danfse"
    assert "CANCELADA" not in _pdf_text(stack.fiscal.get(url, headers=world["ha"]).content)
    job_id = _bundle(stack, world, {"format": "PDF", "document_ids": [d["id"]]}).json()["job_id"]
    assert _run(stack, world, job_id)["download_available"] is True
    _ingest(stack, world["a"], world["a_cnpj"], _event(k, "101101", "<xDesc>Cancelamento de NFS-e</xDesc><cMotivo>1</cMotivo>"
                                                      "<xMotivo>Erro</xMotivo>", world["a_cnpj"]))
    # PDF individual: gerado a cada pedido (sem cache) → já sai com CANCELADA
    r = stack.fiscal.get(url, headers=world["ha"])
    assert "CANCELADA" in _pdf_text(r.content) and r.headers["cache-control"] == "no-store"
    # ZIP gerado antes do evento: bloqueado para não distribuir PDF antigo
    job = stack.fiscal.get(f"{F}/{world['a']}/jobs/{job_id}", headers=world["ha"]).json()
    assert job["stale_pdf_documents"] == 1 and job["download_available"] is False and "pdf_watermarks" not in job["result"]
    r = stack.fiscal.get(f"{F}/{world['a']}/jobs/{job_id}/download", headers=world["ha"])
    assert r.status_code == 409 and r.json()["code"] == "BUNDLE_STALE"
    assert stack.fiscal.post(f"{F}/{world['a']}/jobs/{job_id}/download-ticket", headers=world["ha"]).status_code == 409


def test_ticket_bound_to_job_company_user_short_lived_and_not_logged(stack, world, caplog):
    import logging

    caplog.set_level(logging.DEBUG)
    toma = cnpj("445556660001")
    ids = []
    for n in (80, 81):
        k, x = _nota(world["a_cnpj"], toma, n)
        _ingest(stack, world["a"], world["a_cnpj"], x)
        ids.append(_docs(stack, world)[k]["id"])
    j1 = _bundle(stack, world, {"format": "XML", "document_ids": [ids[0]]}).json()["job_id"]
    j2 = _bundle(stack, world, {"format": "XML", "document_ids": [ids[1]]}).json()["job_id"]
    _run(stack, world, j1)
    _run(stack, world, j2)
    t1 = stack.fiscal.post(f"{F}/{world['a']}/jobs/{j1}/download-ticket", headers=world["ha"]).json()["ticket"]
    # ticket do job 1 não baixa o arquivo do job 2
    r = stack.fiscal.post(f"{F}/{world['a']}/jobs/{j2}/download-file", data={"ticket": t1})
    assert r.status_code == 403 and t1 not in r.text
    # validade curta (60 s)
    t2 = stack.fiscal.post(f"{F}/{world['a']}/jobs/{j1}/download-ticket", headers=world["ha"]).json()["ticket"]
    stack.clock.advance(seconds=61)
    assert stack.fiscal.post(f"{F}/{world['a']}/jobs/{j1}/download-file", data={"ticket": t2}).status_code == 403
    # usuário desativado depois de emitir o ticket: recusado na revalidação com o core
    from core_api.models import User

    t3 = stack.fiscal.post(f"{F}/{world['a']}/jobs/{j1}/download-ticket", headers=world["ha"]).json()["ticket"]
    with stack.core_sessions() as s:
        s.get(User, world["ua"]).is_active = False
        s.commit()
    r = stack.fiscal.post(f"{F}/{world['a']}/jobs/{j1}/download-file", data={"ticket": t3})
    assert r.status_code == 403 and t3 not in r.text
    # entrada inválida não ecoa o valor
    bad = "x" * 30 + "TICKETSECRETO"
    r = stack.fiscal.post(f"{F}/{world['a']}/jobs/{j1}/download-file", data={"ticket": bad + "y" * 80})
    assert r.status_code == 422 and "TICKETSECRETO" not in r.text
    logs = caplog.text
    for t in (t1, t2, t3, "TICKETSECRETO"):
        assert t not in logs
    from ctrl_fiscal.models import DownloadTicket

    with stack.fiscal_sessions() as s:  # só o hash é guardado
        assert all(len(x.id) == 64 and x.id not in (t1, t2, t3) for x in s.query(DownloadTicket).all())


def test_code1_presentation_in_filters_summary_and_export(stack, world):
    from openpyxl import load_workbook

    k, x = _nota(world["a_cnpj"], cnpj("445556660001"), 90,
                 fed={"tpRetPisCofins": "1", "vPis": "6.50", "vCofins": "30.00", "vRetCSLL": "10.00"}, v_total_ret="10.00")
    _ingest(stack, world["a"], world["a_cnpj"], x)
    d = _docs(stack, world, retention_type="CONTRIB")[k]  # contribuições retidas pela modalidade 1
    s = d["retention_summary"]
    assert s["contributions"]["amount"] == "46.50" and s["contributions"]["code"] == "1"
    assert s["informed_not_withheld"] == {"pis": None, "cofins": None}  # não tratados como "informados, não retidos"
    assert s["total_withheld"] == "10.00"  # total da consulta = vTotalRet oficial do XML (não recalculado)
    r = stack.fiscal.post(f"{F}/{world['a']}/exports", headers=world["ha"], json={"format": "xlsx", "query": k})
    job_id = r.json()["job_id"]
    stack.worker().run_jobs()
    xlsx = stack.fiscal.get(f"{F}/{world['a']}/jobs/{job_id}/download", headers=world["ha"]).content
    ws = load_workbook(io.BytesIO(xlsx)).active
    rows = list(ws.iter_rows(values_only=True))
    head, row = rows[0], dict(zip(rows[0], rows[1]))
    assert "PIS" not in head and "CSLL" not in head  # colunas antigas enganosas removidas
    assert str(row["Contribuições retidas PIS/COFINS/CSLL (consolidado; NT 008 v1.02)"]) in ("46.5", "46.50")
    assert row["Modalidade tpRetPisCofins"].startswith("1 - PIS/COFINS retidos")
    assert str(row["vRetCSLL (valor do XML)"]) in ("10", "10.0", "10.00")
    assert str(row["Total retido (vTotalRet)"]) in ("10", "10.0", "10.00")


# ---------------------------------------------------------------- fontes: fornecimento local com conferência de licença
def test_font_provider_local_dir_with_embedding_check(tmp_path, world):
    import shutil
    import struct

    from ctrl_fiscal.danfse import gerar
    from ctrl_fiscal.danfse.fontes import ASSETS, fs_type, resolve

    assert resolve("").origem == "SUBSTITUTA_LIBERATION"  # sem diretório configurado: substituta + desvio explícito
    assert "faltam" in resolve(str(tmp_path)).detalhe
    # Diretório local "licenciado" (cópias da Liberation com os nomes esperados; nenhuma fonte proprietária no teste)
    for name, src in (("arial.ttf", "LiberationSans-Regular.ttf"), ("ARIALBD.TTF", "LiberationSans-Bold.ttf"),
                      ("micross.ttf", "LiberationSans-Regular.ttf")):
        shutil.copy(ASSETS / src, tmp_path / name)
    fs = resolve(str(tmp_path))
    assert fs.origem == "LOCAL_LICENCIADA" and "fsType=" in fs.detalhe
    k, x = _nota(world["a_cnpj"] if isinstance(world, dict) else cnpj("112223330001"), cnpj("445556660001"), 95)
    pdf = gerar(x, None, fs)
    assert "fontes=LOCAL_LICENCIADA" in (PdfReader(io.BytesIO(pdf)).metadata.get("/Keywords") or "")
    # Fonte cuja licença proíbe incorporação (fsType Restricted) é recusada → volta à substituta, com o motivo
    data = bytearray((tmp_path / "micross.ttf").read_bytes())
    n = struct.unpack(">H", data[4:6])[0]
    for i in range(n):
        rec = 12 + 16 * i
        if data[rec:rec + 4] == b"OS/2":
            off = struct.unpack(">I", data[rec + 8:rec + 12])[0]
            data[off + 8:off + 10] = struct.pack(">H", 0x0002)
    (tmp_path / "micross.ttf").write_bytes(bytes(data))
    assert fs_type(tmp_path / "micross.ttf") == 0x0002
    fs2 = resolve(str(tmp_path))
    assert fs2.origem == "SUBSTITUTA_LIBERATION" and "incorporação proibida" in fs2.detalhe
