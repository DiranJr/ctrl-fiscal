"""Extrai do Anexo I oficial (Produção Restrita) o leiaute e as regras de negócio dos campos da DPS (infDPS) para
docs/oficial/regras-dps-anexo-i.json, base da matriz de conformidade.

Uso: .venv\\Scripts\\python -I scripts\\extrair_regras_anexo_i.py

Linhas de continuação (sem caminho/campo) herdam o caminho/campo da linha anterior. Mantém: caminho, campo,
ocorrência, tamanho, tipo, descrição/notas (leiaute) e regra, aplicação, efeito, código e nível (RN).
"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SRC = next((ROOT / "docs/oficial/xsd/producao-restrita").glob("anexo_i-*.xlsx"))
OUT = ROOT / "docs/oficial/regras-dps-anexo-i.json"
PREFIX = "NFSe/infNFSe/DPS/infDPS/"


def _s(v) -> str:  # noqa: ANN001
    return "" if v is None else str(v).replace("\n", " ").strip()


def main() -> int:
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    leiaute: dict[str, dict] = {}
    for r in wb["LEIAUTE DPS_NFS-e "].iter_rows(min_row=2, values_only=True):
        path, campo = _s(r[1]), _s(r[2])
        if path.startswith(PREFIX) or path == "NFSe/infNFSe/DPS/":
            leiaute[f"{path}{campo}"] = {"tipo": _s(r[3]), "formato": _s(r[4]), "ocorrencia": _s(r[5]), "tamanho": _s(r[6]),
                                        "descricao": _s(r[7])[:400], "notas": _s(r[8])[:600]}
    regras: list[dict] = []
    last_path = last_campo = ""
    for r in wb["RN DPS_NFS-e"].iter_rows(min_row=4, values_only=True):
        if not r or len(r) < 9:
            continue
        path, campo = _s(r[1]), _s(r[2])
        if path and path != "-":
            last_path = path
        if campo and campo != "-":
            last_campo = campo
        texto, aplic, efeito, codigo, msg = _s(r[3]), _s(r[5]), _s(r[6]), _s(r[7]), _s(r[8])
        if not last_path.startswith(PREFIX) or not codigo or codigo == "-":
            continue
        regras.append({"linha": _s(r[0]), "elemento": f"{last_path}{last_campo}", "regra": texto[:700], "aplicacao": aplic,
                       "efeito": efeito, "codigo": codigo, "mensagem": msg[:300], "nivel": _s(r[9]) if len(r) > 9 else "",
                       "emissor_proprio": _s(r[10]) if len(r) > 10 else ""})
    OUT.write_text(json.dumps({"fonte": SRC.name, "fonte_sha256": hashlib.sha256(SRC.read_bytes()).hexdigest(),
                               "leiaute": leiaute, "regras": regras}, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"leiaute: {len(leiaute)} elementos de infDPS; regras com código: {len(regras)} -> {OUT.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
