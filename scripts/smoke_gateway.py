"""Smoke test de integração LOCAL via gateway (stack Compose em execução, ambiente mock).

Uso: .venv\\Scripts\\python scripts\\smoke_gateway.py [--base http://127.0.0.1:8098]
Roda no AMBIENTE TEMPORÁRIO (scripts/ambiente_descartavel.py), com a conta fictícia operador.a (senha em
secrets-descartavel/; não imprime). Não contata o governo e não usa o ambiente operacional.
"""
from __future__ import annotations

import argparse
import io
import os
import base64
import gzip
import sys
import time
import uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path

import httpx

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "services" / "nfse-api"))
sys.path.insert(0, str(ROOT / "scripts"))


def check(label: str, cond: bool, extra: str = "") -> None:
    print(f"[{'PASS' if cond else 'FAIL'}] {label}{(' — ' + extra) if extra else ''}")
    if not cond:
        raise SystemExit(1)


def wait(fn, timeout=90, every=2.0):  # noqa: ANN001, ANN201
    end = time.time() + timeout
    while time.time() < end:
        v = fn()
        if v:
            return v
        time.sleep(every)
    return None


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", default="http://127.0.0.1:8098")
    ap.add_argument("--email", default=os.environ.get("CTRL_TEST_EMAIL", "operador.a@descartavel.example.com"))
    ap.add_argument("--password-file", default=os.environ.get("CTRL_TEST_PASSWORD_FILE", "secrets-descartavel/senha-operador.a.txt"))
    a = ap.parse_args()
    origin = a.base
    password = (ROOT / a.password_file).read_text(encoding="utf-8").strip()
    c = httpx.Client(base_url=a.base, timeout=30)

    r = c.post("/api/core/v1/auth/login", json={"email": a.email, "password": password})
    check("login sem Origin recusado (403)", r.status_code == 403)
    r = c.post("/api/core/v1/auth/login", json={"email": a.email, "password": password}, headers={"Origin": origin})
    check("login (200) + cookie HttpOnly/SameSite=Strict", r.status_code == 200 and "httponly" in r.headers.get("set-cookie", "").lower()
          and "samesite=strict" in r.headers.get("set-cookie", "").lower(), r.headers.get("set-cookie", "").split(";")[0].split("=")[0])
    access = r.json()["access_token"]
    H = {"Authorization": "Bearer " + access}
    csrf = c.get("/api/core/v1/auth/csrf").json()["csrf_token"]
    check("GET /auth/csrf com cookie", bool(csrf))
    check("refresh sem CSRF recusado (403)", c.post("/api/core/v1/auth/refresh", headers={"Origin": origin}).status_code == 403)
    r = c.post("/api/core/v1/auth/refresh", headers={"Origin": origin, "X-CSRF-Token": csrf})
    check("refresh com rotação (200)", r.status_code == 200)
    H = {"Authorization": "Bearer " + r.json()["access_token"]}
    # Recarga da página: novo contexto sem token/CSRF em memória, apenas o cookie HttpOnly.
    reload_ = httpx.Client(base_url=a.base, timeout=30, cookies=c.cookies)
    check("recarga: refresh sem CSRF => 403", reload_.post("/api/core/v1/auth/refresh", headers={"Origin": origin}).status_code == 403)
    csrf2 = reload_.get("/api/core/v1/auth/csrf").json()["csrf_token"]
    r2 = reload_.post("/api/core/v1/auth/refresh", headers={"Origin": origin, "X-CSRF-Token": csrf2})
    check("recarga: GET /auth/csrf + refresh => sessão recuperada", r2.status_code == 200 and csrf2 == csrf)
    c.cookies = reload_.cookies
    H = {"Authorization": "Bearer " + r2.json()["access_token"]}
    me = c.get("/api/core/v1/auth/me", headers=H).json()
    check("GET /auth/me com empresas", len(me["companies"]) >= 1, f"{len(me['companies'])} empresas")
    company = me["companies"][0]
    cid = company["id"]
    F = f"/api/fiscal/v1/companies/{cid}"
    check("empresa não associada => 403", c.get(f"/api/fiscal/v1/companies/{uuid.uuid4()}/documents", headers=H).status_code == 403)
    check("sem token => 401 problem+json", c.get(f"{F}/documents").headers["content-type"].startswith("application/problem+json"))
    dash = c.get(f"{F}/dashboard", headers=H)
    check("dashboard", dash.status_code == 200 and dash.json()["simulated"] is True, f"env={dash.json()['environment']}")

    # Sincronização ADN (mock): o worker executa o ciclo inicial sozinho.
    st = wait(lambda: (lambda s: s if s["status"] in ("OK", "ERROR", "NOT_CONFIGURED", "BLOCKED") else None)(
        c.get(f"{F}/sync/status", headers=H).json()))
    if st is not None and st["status"] == "BLOCKED":
        # Sprint 02: mock desligado em execução; o ciclo não consulta nada e registra o bloqueio.
        check("sync ADN (mock desligado) => BLOCKED/MOCK_DISABLED, sem fallback",
              (st.get("last_error") or {}).get("code") == "MOCK_DISABLED" and st["interval_hours"] == 12,
              f"status={st['status']} erro={st.get('last_error')} interval={st['interval_hours']}h")
    else:
        check("sync ADN ciclo inicial pelo worker", st is not None and st["status"] == "OK",
              f"status={st and st['status']} next_run_at={st and st['next_run_at']} interval={st and st['interval_hours']}h")
    run = c.post(f"{F}/sync/run", headers={**H, "Idempotency-Key": uuid.uuid4().hex})
    check("sync/run => 202 job (sem consulta no request)", run.status_code == 202, run.json()["status"])

    # Cadastro de clientes (payload canônico; payload da tela atual é recusado com 422)
    base12 = str(int(time.time() * 1000))[-12:]
    from ctrl_common_cnpj import cnpj_full
    newdoc = cnpj_full(base12)
    bad = c.post(f"/api/core/v1/companies/{cid}/customers", headers=H, json={"name": "UI", "cnpj_cpf": newdoc})
    check("cliente com payload legado (cnpj_cpf) => 422", bad.status_code == 422)
    cc = c.post(f"/api/core/v1/companies/{cid}/customers", headers=H,
                json={"name": "Cliente Smoke", "document_type": "CNPJ", "document": newdoc})
    check("cadastro de cliente canônico (201) e duplicado (409)", cc.status_code == 201 and c.post(
        f"/api/core/v1/companies/{cid}/customers", headers=H,
        json={"name": "Dup", "document_type": "CNPJ", "document": newdoc}).status_code == 409)
    pc = c.patch(f"/api/core/v1/companies/{cid}/customers/{cc.json()['id']}", headers=H, json={"name": "Cliente Smoke 2", "expected_version": 1})
    check("edição de cliente com versão", pc.status_code == 200 and pc.json()["version"] == 2)

    # Usar última nota do cliente -> rascunho -> lote -> worker (mock)
    customers = c.get(f"/api/core/v1/companies/{cid}/customers", headers=H).json()["items"]
    cust = next(x for x in customers if x["document"] == "44555666000181")
    last = c.get(f"{F}/customers/{cust['id']}/last-issued-document", headers=H)
    check("última nota autorizada do cliente (ADN mock)", last.status_code == 200, last.json().get("number", ""))
    dr = c.post(f"{F}/drafts/from-document", headers={**H, "Idempotency-Key": uuid.uuid4().hex}, json={"source_document_id": last.json()["id"]}).json()
    check("rascunho por referência sem IDs fiscais", dr["status"] == "DRAFT" and "access_key" not in dr and "number" not in dr)
    dr = c.patch(f"{F}/drafts/{dr['id']}", headers=H, json={"expected_version": dr["version"], "service": {"national_code": "010101"},
                                                            "tax_settings": {"issqn_taxation": "TAXABLE", "iss_retention": "NOT_RETAINED"}}).json()
    val = c.post(f"{F}/drafts/{dr['id']}/validate", headers=H).json()
    check("validação local", val["valid"] is True, ",".join(i["code"] for i in val["issues"]))
    at = (datetime.now(timezone.utc) + timedelta(seconds=5)).isoformat().replace("+00:00", "Z")  # formato toISOString()
    b = c.post(f"{F}/batches", headers={**H, "Idempotency-Key": uuid.uuid4().hex},
               json={"drafts": [{"id": dr["id"], "expected_version": dr["version"]}], "scheduled_at": at, "timezone": company["timezone"]})
    check("lote agendado com 'Z' (202, resposta em UTC)", b.status_code == 202 and b.json()["scheduled_at"].endswith("+00:00"), b.text[:120])
    bd = c.get(f"{F}/batches/{b.json()['batch_id']}", headers=H).json()
    check("fuso validado e hora local registrados", bd["timezone"] == company["timezone"] and bd["scheduled_local"].endswith("-03:00")
          and bd["items"][0]["competence"] == dr["competence"], f"tz={bd['timezone']} local={bd['scheduled_local']} comp={bd['items'][0]['competence']}")
    wrong = c.post(f"{F}/batches", headers={**H, "Idempotency-Key": uuid.uuid4().hex},
                   json={"drafts": [{"id": dr["id"], "expected_version": dr["version"]}], "scheduled_at": at, "timezone": "UTC"})
    check("fuso diferente do da empresa => 422", wrong.status_code == 422 and wrong.json()["code"] == "TIMEZONE_MISMATCH")
    integ = c.get(f"/api/core/v1/companies/{cid}/integration-status", headers=H).json()
    mock_off = any(x.get("code") == "MOCK_DISABLED" for x in integ.get("blockers", []))
    if mock_off:
        # Sprint 02: mock desligado em execução (sem fallback). O item não é transmitido e fica agendado.
        check("integração mock => BLOCKED/MOCK_DISABLED", integ["status"] == "BLOCKED" and integ["capabilities"]["emission"] is False)
        item = wait(lambda: (lambda d: d["items"][0] if (d["items"][0].get("last_error") or {}).get("code") else None)(
            c.get(f"{F}/batches/{b.json()['batch_id']}", headers=H).json()))
        check("worker NÃO transmitiu: item SCHEDULED com MOCK_DISABLED (sem fallback)",
              item is not None and item["status"] == "SCHEDULED" and item["last_error"]["code"] == "MOCK_DISABLED",
              item and f"{item['status']} {item.get('last_error')}")
        wd = c.post(f"{F}/batches/{b.json()['batch_id']}/cancel-pending", headers={**H, "Idempotency-Key": uuid.uuid4().hex})
        check("lote de teste retirado (cancel-pending)", wd.status_code == 200, wd.text[:120])
    else:
        item = wait(lambda: (lambda d: d["items"][0] if d["items"][0]["status"] in ("AUTHORIZED", "REJECTED") else None)(
            c.get(f"{F}/batches/{b.json()['batch_id']}", headers=H).json()))
        check("worker transmitiu (SIMULADA)", item is not None and item["status"] == "AUTHORIZED" and item["result"]["simulated"],
              item and item["result"]["label"])

    # Importação XML (mock) + conferência + download autenticado
    from nfse_api.providers.mock_data import cnpj_dv, mock_access_key, nfse_xml

    key = mock_access_key(cnpj_dv("700000000001"), 9001, datetime.now(timezone.utc), uuid.uuid4().hex)
    xml = nfse_xml(key=key, number=9001, prest=(cnpj_dv("700000000001"), "Fornecedor Importado"), toma=(company["cnpj"], "Empresa"),
                   interm=None, description="Importação smoke", v_serv="100.00", issued_at=datetime.now(timezone(timedelta(hours=-3))),
                   competence_day="2026-10-01")
    key0 = mock_access_key(cnpj_dv("700000000001"), 9000, datetime.now(timezone.utc), uuid.uuid4().hex)
    xml0 = nfse_xml(key=key0, number=9000, prest=(cnpj_dv("700000000001"), "Fornecedor Importado"), toma=(company["cnpj"], "Empresa"),
                    interm=None, description="Importação smoke OK", v_serv="90.00", issued_at=datetime.now(timezone(timedelta(hours=-3))),
                    competence_day="2026-10-01")
    r0 = c.post(f"{F}/documents/import", headers=H, files=[("files", ("so-valido.xml", xml0, "application/xml"))])
    job0 = wait(lambda: (lambda j: j if j["status"] not in ("QUEUED", "RUNNING") else None)(
        c.get(f"{F}/jobs/{r0.json()['job_id']}", headers=H).json()))
    check("import só válido => job SUCCEEDED", job0 and job0["status"] == "SUCCEEDED", job0 and str(job0["result"]))
    evil = b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]><NFSe>&e;</NFSe>'
    bad_id = b'<?xml version="1.0"?><NFSe xmlns="http://www.sped.fazenda.gov.br/nfse"><infNFSe Id="NFS123"/></NFSe>'
    r = c.post(f"{F}/documents/import", headers=H, files=[("files", ("nota.xml", xml, "application/xml")),
                                                         ("files", ("invalida.xml", bad_id, "application/xml"))])
    check("import 202", r.status_code == 202)
    job = wait(lambda: (lambda j: j if j["status"] not in ("QUEUED", "RUNNING") else None)(
        c.get(f"{F}/jobs/{r.json()['job_id']}", headers=H).json()))
    found = c.get(f"{F}/documents", headers=H, params={"query": key}).json()
    check("válido + inválido => PARTIAL, válido persistido, erro do inválido",
          job and job["status"] == "PARTIAL" and found["total"] == 1
          and any(e["file"] == "invalida.xml" for e in job["errors"]),
          job and f"{job['result']} erros={[(e['file'], e['code']) for e in job['errors']]}")
    r = c.post(f"{F}/documents/import", headers=H, files=[("files", ("xxe.xml", evil, "application/xml")),
                                                         ("files", ("invalida2.xml", bad_id, "application/xml"))])
    jobf = wait(lambda: (lambda j: j if j["status"] not in ("QUEUED", "RUNNING") else None)(
        c.get(f"{F}/jobs/{r.json()['job_id']}", headers=H).json()))
    check("importação totalmente inválida => FAILED (XXE bloqueado)", jobf and jobf["status"] == "FAILED"
          and "XML_FORBIDDEN_CONSTRUCT" in {e["code"] for e in jobf["errors"]}, jobf and str([e["code"] for e in jobf["errors"]]))
    doc = found["items"][0]
    x = c.get(f"{F}/documents/{doc['id']}/xml", headers=H)
    check("XML download autenticado = original", x.status_code == 200 and x.content == xml)
    check("XML sem token => 401", c.get(f"{F}/documents/{doc['id']}/xml").status_code == 401)
    rv = c.patch(f"{F}/documents/{doc['id']}/review", headers=H, json={"review_status": "REVIEWED", "expected_version": doc["version"]})
    check("conferência registrada", rv.status_code == 200 and rv.json()["reviewed_by"] is not None)

    # Exportação XLSX assíncrona: abre o arquivo e confere dados, filtro e isolamento por empresa
    def export_rows(company_id, body):
        e = c.post(f"/api/fiscal/v1/companies/{company_id}/exports", headers=H, json=body)
        j = wait(lambda: (lambda jj: jj if jj["status"] not in ("QUEUED", "RUNNING") else None)(
            c.get(f"/api/fiscal/v1/companies/{company_id}/jobs/{e.json()['job_id']}", headers=H).json()))
        d = c.get(f"/api/fiscal/v1/companies/{company_id}/jobs/{e.json()['job_id']}/download", headers=H)
        from openpyxl import load_workbook
        rows = list(load_workbook(io.BytesIO(d.content), read_only=True).active.iter_rows(values_only=True))
        return {**j, "_job_id": e.json()["job_id"], "_bytes": len(d.content)}, rows

    jobx, rows = export_rows(cid, {"format": "xlsx", "direction": "RECEIVED", "page": 1, "page_size": 25})
    header, data = rows[0], rows[1:]
    total_api = c.get(f"{F}/documents", headers=H, params={"direction": "RECEIVED", "page_size": 1}).json()["total"]
    col = {name: i for i, name in enumerate(header)}
    mine = [row for row in data if row[col["Chave de acesso"]] == key]
    check("XLSX aberto: cabeçalho, todas as linhas do filtro (sem paginação) e só 'Recebida'",
          jobx["status"] == "SUCCEEDED" and len(data) == total_api == jobx["result"]["rows"]
          and all(row[col["Direção"]] == "Recebida" for row in data),
          f"linhas={len(data)} api_total={total_api} colunas={len(header)}")
    check("XLSX contém o documento importado com valores corretos",
          len(mine) == 1 and float(mine[0][col["Valor do serviço"]]) == 100.0 and mine[0][col["Conferência"]] == "REVIEWED",
          mine and f"nº {mine[0][col['Número']]} valor={mine[0][col['Valor do serviço']]} conferência={mine[0][col['Conferência']]}")
    other = next(co for co in me["companies"] if co["id"] != cid)
    jobb, rows_b = export_rows(other["id"], {"format": "xlsx"})
    keys_b = {row[col["Chave de acesso"]] for row in rows_b[1:]}
    keys_a = {row[col["Chave de acesso"]] for row in data}
    check("isolamento: export da outra empresa não contém documentos desta", jobb["status"] == "SUCCEEDED" and key not in keys_b
          and not (keys_a & keys_b), f"{other['name']}: {len(rows_b) - 1} linhas; interseção={len(keys_a & keys_b)}")
    check("download de export sem token => 401", c.get(f"{F}/jobs/{jobx['_job_id']}/download").status_code == 401)
    check("download de export pela rota de outra empresa => 404", c.get(
        f"/api/fiscal/v1/companies/{other['id']}/jobs/{jobx['_job_id']}/download", headers=H).status_code == 404)

    old_cookies = httpx.Cookies(c.cookies)
    r = c.post("/api/core/v1/auth/logout", headers={"Origin": origin, "X-CSRF-Token": csrf})
    check("logout 204", r.status_code == 204)
    check("token após logout => 401 (core e fiscal)", c.get("/api/core/v1/auth/me", headers=H).status_code == 401
          and c.get(f"{F}/documents", headers=H).status_code == 401)
    replay = httpx.Client(base_url=a.base, timeout=30, cookies=old_cookies)
    rr = replay.post("/api/core/v1/auth/refresh", headers={"Origin": origin, "X-CSRF-Token": csrf})
    check("reuso do cookie de refresh revogado => 401 (refresh e csrf)", rr.status_code == 401
          and replay.get("/api/core/v1/auth/csrf").status_code == 401, f"refresh={rr.status_code} {rr.json().get('code')}")
    frontend_bundle_report(a.base)
    print("SMOKE OK (ambiente mock; nenhuma chamada governamental)")


def frontend_bundle_report(base: str) -> None:
    """Informativo: o frontend servido pelo gateway usa o cliente HTTP real ou mocks no navegador?"""
    import re

    markers = ("22.333.444/0001-55", "33.444.555/0001-66")  # fixtures de packages/api-client/src/mock.ts
    c = httpx.Client(base_url=base, timeout=30)
    for label, page in (("painel /fiscal/", "/fiscal/"), ("web /login", "/login")):
        html = c.get(page).text
        assets = set(re.findall(r'(?:src|href)="(/[^"]+\.js)"', html))
        found = set()
        for asset in assets:
            body = c.get(asset).text
            found |= {m for m in markers if m in body}
        print(f"[INFO] {label}: {len(assets)} bundles JS; fixtures de mock presentes no bundle: {sorted(found) or 'nenhuma'} "
              "(presença não prova uso: o módulo mock é exportado pelo pacote; uso real conferido no código e por HTTP)")
    credential_report(base)



def credential_report(base: str) -> None:
    """Informativo de segurança: o login servido traz credencial pré-preenchida? (sem imprimir valores)"""
    import re

    src = (ROOT / "apps" / "web" / "src" / "pages" / "login.tsx").read_text(encoding="utf-8")
    em = re.search(r"\[email, setEmail\] = useState\('([^']*)'\)", src)
    pw = re.search(r"\[password, setPassword\] = useState\('([^']*)'\)", src)
    em, pw = (em.group(1) if em else ""), (pw.group(1) if pw else "")
    c = httpx.Client(base_url=base, timeout=30)
    bodies = [c.get(j).text for j in set(re.findall(r'/_next/static/[^"]+\.js', c.get("/login").text))]
    print(f"[INFO] login.tsx pré-preenche e-mail={bool(em)} senha={bool(pw)}; no bundle servido: "
          f"e-mail={any(em and em in b for b in bodies)} senha={any(pw and pw in b for b in bodies)}")
    if pw:
        r = c.post("/api/core/v1/auth/login", json={"email": em, "password": pw}, headers={"Origin": base})
        print(f"[INFO] credencial embutida autentica no backend: {r.status_code == 200} (status {r.status_code})")


if __name__ == "__main__":
    main()
