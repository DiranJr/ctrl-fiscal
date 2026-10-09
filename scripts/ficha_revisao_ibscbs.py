"""Gera a ficha curta de revisão contábil da 1ª emissão em Produção Restrita: serviço, descrição, local da prestação,
tomador e os códigos IBS/CBS pendentes (com os valores da nota de referência, SEM preenchê-los).

Uso: .venv\\Scripts\\python -I scripts\\ficha_revisao_ibscbs.py

Imprime a ficha com o tomador MASCARADO (para relatórios/chat). Com --completa, imprime o tomador sem máscara
(dados reais: só no terminal local; nenhum arquivo é gravado).
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OFIC = ROOT / "docs/oficial/xsd/producao-restrita"


def _lookup(xlsx: Path, sheet: str | None, code: str, code_cols: tuple[int, ...], text_col: int) -> str | None:
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    for row in ws.iter_rows(values_only=True):
        for c in code_cols:
            if c < len(row) and row[c] is not None and str(row[c]).replace(".", "").strip().lstrip("0") == code.lstrip("0"):
                return str(row[text_col]).replace("\n", " ").strip() if row[text_col] else None
    return None


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")  # saída no terminal (Windows usa cp1252 por padrão)
    nota = json.loads((ROOT / "data/homologacao/nota-teste.json").read_text(encoding="utf-8"))
    n, tom, svc = nota["nota"], nota["tomador"], nota["nota"]["service"]
    ref = (nota.get("_revisao_contabil_ibs_cbs") or {}).get("valores_da_referencia") or {}
    anexo_i = next(OFIC.glob("anexo_i-*.xlsx"))
    anexo_b = next(OFIC.glob("anexo_b-*.xlsx"))
    anexo_c = next(OFIC.glob("anexo_c-*.xlsx"))
    serv_txt = _lookup(anexo_i, "MUN.INCID_INFO.SERV.", svc["national_code"], (0,), 1)
    nbs_txt = None
    wb = openpyxl.load_workbook(anexo_b, read_only=True, data_only=True)
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if any(c.replace(".", "").strip() == svc.get("nbs_code") for c in cells):
                nbs_txt = max(cells, key=len)[:160]
                break
        if nbs_txt:
            break
    indop_txt = None
    grupo = None  # cabeçalho do inciso (o Anexo C repete só o código nas linhas de continuação)
    for row in openpyxl.load_workbook(anexo_c, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True):
        cells = [str(c).replace("\n", " ").strip() for c in row if c is not None]
        if cells and cells[0].startswith("Inc."):
            grupo = f"{cells[0]}: {cells[1]} (local da operação: {cells[2]})" if len(cells) > 2 else cells[0]
        if (ref.get("cIndOp") or "") in cells:
            i = cells.index(ref["cIndOp"])
            indop_txt = f"{grupo or '?'}; local do fornecimento: {cells[i + 1] if i + 1 < len(cells) else '-'}"
            break
    end = tom.get("address") or {}

    def ficha(mask: bool) -> str:
        doc = tom["document"]
        doc_s = (doc[:2] + "*" * (len(doc) - 6) + doc[-4:]) if mask else doc
        nome = (tom["name"][:3] + "…") if mask else tom["name"]
        desc = (svc["description"][:40] + "…") if mask else svc["description"]
        return "\n".join([
            "# Ficha de revisão — 1ª emissão em Produção Restrita (IBS/CBS pendente)",
            "",
            f"- **Serviço (cTribNac):** {svc['national_code']} — {serv_txt or '(texto não localizado no Anexo I)'}",
            f"- **NBS:** {svc.get('nbs_code')} — {nbs_txt or '(texto não localizado no Anexo B)'}",
            f"- **Descrição (xDescServ):** {desc}",
            f"- **Local da prestação / incidência ISSQN:** {svc['location_code']} (regra LP do Anexo I)",
            f"- **Competência / data da prestação:** {n['competence']} / {n['service_date']}",
            f"- **Valor do serviço:** {n['amount']}  ·  ISSQN retido: {n['tax_settings']['iss_retention']}  ·  pAliq: {n['tax_settings']['iss_rate']}",
            f"- **Tomador:** {nome} — CNPJ {doc_s} — município {tom.get('municipality_code')}"
            + ("" if mask else f" — {end.get('street')}, {end.get('number')} — {end.get('district')} — CEP {tom.get('postal_code')}"),
            "",
            "## Códigos IBS/CBS pendentes (não preenchidos)",
            "",
            "| Campo | Valor na nota de referência | Revisar |",
            "|---|---|---|",
            f"| cIndOp | {ref.get('cIndOp')} | {indop_txt or '-'} |",
            f"| CST IBS/CBS | {ref.get('CST')} | tabela CST da Reforma (não disponível localmente) |",
            f"| cClassTrib | {ref.get('cClassTrib')} | tabela cClassTrib (pode exigir gTribRegular/gDif — E1583/E1584) |",
            f"| indFinal | {ref.get('indFinal')} | uso/consumo pessoal (art. 57): 0 não, 1 sim |",
            "| indDest | 0 (fixo) | tomador = destinatário; outro destinatário não é suportado |",
            "",
            "Decisão pedida: confirmar ou corrigir cIndOp, CST, cClassTrib e indFinal para ESTE serviço e ESTA data.",
        ])
    sys.stdout.write(ficha("--completa" not in sys.argv) + "\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
